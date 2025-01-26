# Used to open and manipulate word documents
from docx import Document
from docx.shared import Pt, RGBColor
from io import BytesIO
from preview import Preview
import os
import subprocess

class CoverLetter():

    def __init__(self, document_path, date, company, city, position, s, destination):

        # User inputs
        self.document = Document(document_path)
        self.date = date
        self.company = company
        self.city = city
        self.position = position
        self.s = s
        self.destinaton = destination

        # Define style
        self.defineDocumentStyles()

        # Create dictionary based on user input - will replace placeholders in doc
        self.replacements = {
            '{DATE}': self.date,
            '{COMPANY}': self.company,
            '{CITY}': self.city,
            '{C_POSITION}': self.position,
            '{POSITION}': self.position,
            '{S}': self.s
        }

    def defineDocumentStyles(self):
        """
        Define the document styles for the cover letter.
        """
        self.reg_style = self.document.styles['CL_Normal']
        font_1 = self.reg_style.font
        font_1.name = 'Calibri'
        font_1.size = Pt(11)

        self.coloured_style = self.document.styles['Coloured']
        font_2 = self.coloured_style.font
        font_2.name = 'Calibri'
        font_2.size = Pt(11)
        font_2.color.rgb = RGBColor(74, 134, 232)
        font_2.bold = True
    
    # TODO: edit colour and font in one placeholder
    def replacePlaceholders(self):
        """
        Replace placeholders in the cover letter and save new cover letter.
        """
        all_paragraphs = self.document.paragraphs

        for key, value in self.replacements.items():
            for paragraph in all_paragraphs:
                if key in paragraph.text:
                    if key == "{C_POSITION}":
                        print(f"{key}")
                        paragraph.text = paragraph.text.replace(key, value)
                        paragraph.style = self.coloured_style
                    else:
                        print(f"{key}")
                        paragraph.text = paragraph.text.replace(key, value)
                        paragraph.style = self.reg_style

        # save as pdf
        file_path = os.path.join(self.destinaton, f"{self.company}_Arianna_Foo_Cover_Letter.docx")
        self.document.save(file_path)
        print(f"Converting from: {file_path}")
        print(f"Saving PDF to: {self.destinaton}")
        self.convert_docx_to_pdf(file_path, self.destinaton)
        
    
    def convert_docx_to_pdf(self, docx_path, output_pdf_path):
        # Call LibreOffice to convert the document
        subprocess.run(['/Applications/LibreOffice.app/Contents/MacOS/soffice', '--headless', '--convert-to', 'pdf', docx_path, '--outdir', output_pdf_path])

        # generate preview
        new_preview = Preview(f"{output_pdf_path}/{self.company}_AriannaFoo_CL.pdf", self.company)
        new_preview.generate_preview()

    """ TODO
    - save job to excel sheet
    - better way to open folder label
    """